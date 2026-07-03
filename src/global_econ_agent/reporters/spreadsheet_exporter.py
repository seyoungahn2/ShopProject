import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from global_econ_agent.models.schemas import DailyReport
from global_econ_agent.utils.labels import category_label, impact_label, region_label


class SpreadsheetExporter:
    """CSV·Excel 형식으로 보고서를보냅니다. 구글 드라이브에 업로드하면 스프레드시트로 열 수 있습니다."""

    def __init__(self, reports_dir: Path):
        self.reports_dir = reports_dir
        self.reports_dir.mkdir(parents=True, exist_ok=True)

    def export_csv(self, report: DailyReport, base_name: str) -> Path:
        """핵심 이슈를 단일 CSV로 저장 (구글 시트 가져오기에 적합)."""
        output_path = self.reports_dir / f"{base_name}.csv"

        with open(output_path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow([
                "보고서 날짜", report.report_date,
            ])
            writer.writerow([
                "경영진 요약", report.executive_summary,
            ])
            writer.writerow([])

            writer.writerow([
                "번호", "한글 제목", "원문 제목", "영향도", "카테고리", "영향 지역",
                "주식 영향", "채권 영향", "금리 전망", "요약",
                "영향 섹터", "관련 기업", "용어 해석",
            ])
            for i, issue in enumerate(report.key_issues, 1):
                glossary_text = "; ".join(
                    f"{g.term}: {g.meaning_ko}" for g in issue.term_glossary
                )
                writer.writerow([
                    i,
                    issue.title_ko,
                    issue.title_original,
                    impact_label(issue.impact_level),
                    ", ".join(category_label(c) for c in issue.categories),
                    ", ".join(region_label(r) for r in issue.regions_affected),
                    issue.stock_impact,
                    issue.bond_impact,
                    issue.rate_outlook,
                    issue.summary_ko,
                    ", ".join(issue.affected_sectors),
                    ", ".join(issue.affected_companies),
                    glossary_text,
                ])

            writer.writerow([])
            writer.writerow(["시나리오 전망"])
            writer.writerow([
                "시나리오", "확률", "설명", "미국", "한국", "일본", "금리",
                "촉발 요인", "주요 리스크", "투자 시사점",
            ])
            for scenario in report.scenarios:
                writer.writerow([
                    scenario.name,
                    scenario.probability,
                    scenario.description,
                    scenario.us_market_outlook,
                    scenario.kr_market_outlook,
                    scenario.jp_market_outlook,
                    scenario.rate_outlook,
                    "; ".join(scenario.triggers),
                    "; ".join(scenario.key_risks),
                    scenario.investment_implications,
                ])

            if report.term_glossary:
                writer.writerow([])
                writer.writerow(["전체 용어 해석집"])
                writer.writerow(["원어", "한글 해석"])
                for g in report.term_glossary:
                    writer.writerow([g.term, g.meaning_ko])

        return output_path

    def export_xlsx(self, report: DailyReport, base_name: str) -> Path:
        """여러 시트로 구성된 Excel 파일 저장."""
        output_path = self.reports_dir / f"{base_name}.xlsx"
        wb = Workbook()

        self._write_summary_sheet(wb.active, report)
        self._write_issues_sheet(wb.create_sheet("핵심이슈"), report)
        self._write_scenarios_sheet(wb.create_sheet("시나리오"), report)
        if report.term_glossary:
            self._write_glossary_sheet(wb.create_sheet("용어해석"), report)

        wb.save(output_path)
        return output_path

    def _write_summary_sheet(self, ws, report: DailyReport) -> None:
        ws.title = "요약"
        header_font = Font(bold=True, size=12)
        header_fill = PatternFill("solid", fgColor="E8F0FE")

        rows = [
            ("보고서 날짜", report.report_date),
            ("생성 시각", report.generated_at.strftime("%Y-%m-%d %H:%M UTC")),
            ("수집 기사", f"{report.articles_collected}건"),
            ("분석 이슈", f"{report.articles_analyzed}건"),
            ("", ""),
            ("경영진 요약", report.executive_summary),
        ]
        for i, (label, value) in enumerate(rows, 1):
            ws.cell(row=i, column=1, value=label).font = header_font
            ws.cell(row=i, column=1).fill = header_fill
            ws.cell(row=i, column=2, value=value).alignment = Alignment(wrap_text=True)

        ws.column_dimensions["A"].width = 16
        ws.column_dimensions["B"].width = 80

    def _write_issues_sheet(self, ws, report: DailyReport) -> None:
        headers = [
            "번호", "한글 제목", "원문 제목", "영향도", "카테고리", "영향 지역",
            "주식 영향", "채권 영향", "금리 전망", "요약",
            "영향 섹터", "관련 기업", "용어 해석",
        ]
        self._write_header_row(ws, headers)

        for i, issue in enumerate(report.key_issues, 1):
            glossary_text = "\n".join(
                f"{g.term}: {g.meaning_ko}" for g in issue.term_glossary
            )
            row = [
                i,
                issue.title_ko,
                issue.title_original,
                impact_label(issue.impact_level),
                ", ".join(category_label(c) for c in issue.categories),
                ", ".join(region_label(r) for r in issue.regions_affected),
                issue.stock_impact,
                issue.bond_impact,
                issue.rate_outlook,
                issue.summary_ko,
                ", ".join(issue.affected_sectors),
                ", ".join(issue.affected_companies),
                glossary_text,
            ]
            self._write_data_row(ws, i + 1, row)

        self._auto_width(ws)

    def _write_scenarios_sheet(self, ws, report: DailyReport) -> None:
        headers = [
            "시나리오", "확률", "설명", "미국 전망", "한국 전망", "일본 전망",
            "금리 전망", "촉발 요인", "주요 리스크", "투자 시사점",
        ]
        self._write_header_row(ws, headers)

        for i, scenario in enumerate(report.scenarios, 1):
            row = [
                scenario.name,
                scenario.probability,
                scenario.description,
                scenario.us_market_outlook,
                scenario.kr_market_outlook,
                scenario.jp_market_outlook,
                scenario.rate_outlook,
                "\n".join(scenario.triggers),
                "\n".join(scenario.key_risks),
                scenario.investment_implications,
            ]
            self._write_data_row(ws, i + 1, row)

        self._auto_width(ws)

    def _write_glossary_sheet(self, ws, report: DailyReport) -> None:
        self._write_header_row(ws, ["원어", "한글 해석"])
        for i, g in enumerate(report.term_glossary, 1):
            self._write_data_row(ws, i + 1, [g.term, g.meaning_ko])
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 60

    def _write_header_row(self, ws, headers: list[str]) -> None:
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="1A73E8")
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

    def _write_data_row(self, ws, row: int, values: list) -> None:
        for col, value in enumerate(values, 1):
            ws.cell(row=row, column=col, value=value).alignment = Alignment(wrap_text=True, vertical="top")

    def _auto_width(self, ws) -> None:
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_len = max(max_len, min(len(str(cell.value)), 50))
            ws.column_dimensions[col_letter].width = max(max_len + 2, 12)
